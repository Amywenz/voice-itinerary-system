"""把 docs/ 下的 Markdown 源稿导出为 docx（不引入 pandoc 依赖）。

用法：python scripts/build_docs.py
输出：
  docs/软件需求规格说明书_SRS_v1.0.docx
  docs/软件设计说明书_SDS_v1.0.docx
  docs/测试用例表.docx
  docs/模块测试记录.docx
  docs/结项报告.docx
  docs/评审自查表.docx
"""
from __future__ import annotations

import os
import re

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
DOCS = os.path.join(ROOT, "docs")

ACCENT = RGBColor(0x0F, 0x76, 0x6E)
INK = RGBColor(0x1F, 0x29, 0x37)
GRAY = RGBColor(0x47, 0x55, 0x69)
CN_FONT = "微软雅黑"
MONO_FONT = "Consolas"

JOBS = [
    ("SRS.md", "软件需求规格说明书_SRS_v1.0.docx", False),
    ("SDS.md", "软件设计说明书_SDS_v1.0.docx", False),
    ("测试用例表.md", "测试用例表.docx", True),
    ("模块测试记录.md", "模块测试记录.docx", False),
    ("结项报告.md", "结项报告.docx", False),
    ("评审自查表.md", "评审自查表.docx", True),
]


def set_run_font(run, name=CN_FONT, size=None, bold=None, color=None, italic=None):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), name)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    if color is not None:
        run.font.color.rgb = color


def add_page_number(section):
    """页脚居中显示“第 X 页”。"""
    paragraph = section.footer.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run("第 ")
    set_run_font(run, size=9, color=GRAY)

    field_run = paragraph.add_run()
    set_run_font(field_run, size=9, color=GRAY)
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = "PAGE"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    for node in (begin, instr, end):
        field_run._r.append(node)

    run2 = paragraph.add_run(" 页")
    set_run_font(run2, size=9, color=GRAY)


def setup_document(landscape: bool) -> Document:
    doc = Document()
    section = doc.sections[0]
    if landscape:
        section.orientation = WD_ORIENT.LANDSCAPE
        section.page_width, section.page_height = Cm(29.7), Cm(21.0)
    else:
        section.orientation = WD_ORIENT.PORTRAIT
        section.page_width, section.page_height = Cm(21.0), Cm(29.7)
    section.top_margin = Cm(2.2)
    section.bottom_margin = Cm(2.2)
    section.left_margin = Cm(2.2)
    section.right_margin = Cm(2.2)
    add_page_number(section)

    normal = doc.styles["Normal"]
    normal.font.name = CN_FONT
    normal.font.size = Pt(10.5)
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), CN_FONT)
    normal.paragraph_format.line_spacing = 1.4
    normal.paragraph_format.space_after = Pt(4)

    sizes = {"Heading 1": 18, "Heading 2": 14.5, "Heading 3": 12, "Heading 4": 11}
    for style_name, size in sizes.items():
        style = doc.styles[style_name]
        style.font.name = CN_FONT
        style.element.rPr.rFonts.set(qn("w:eastAsia"), CN_FONT)
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = ACCENT if style_name == "Heading 1" else INK
        style.paragraph_format.space_before = Pt(12 if style_name == "Heading 1" else 8)
        style.paragraph_format.space_after = Pt(6)

    return doc


INLINE = re.compile(r"(\*\*.+?\*\*|`[^`]+`|\[[^\]]+\]\([^)]+\))")


def add_inline(paragraph, text: str, base_size=10.5, color=INK):
    """支持 **粗体**、`代码`、[文字](链接) 三种行内标记。"""
    for chunk in INLINE.split(text):
        if not chunk:
            continue
        if chunk.startswith("**") and chunk.endswith("**"):
            run = paragraph.add_run(chunk[2:-2])
            set_run_font(run, size=base_size, bold=True, color=color)
        elif chunk.startswith("`") and chunk.endswith("`"):
            run = paragraph.add_run(chunk[1:-1])
            set_run_font(run, name=MONO_FONT, size=base_size - 0.5, color=RGBColor(0xB4, 0x1C, 0x1C))
        elif chunk.startswith("[") and "](" in chunk:
            label = chunk[1:chunk.index("]")]
            run = paragraph.add_run(label)
            set_run_font(run, size=base_size, color=ACCENT)
        else:
            run = paragraph.add_run(chunk)
            set_run_font(run, size=base_size, color=color)


def split_row(line: str) -> list[str]:
    """按未转义的 | 拆分表格行，并还原 \\|。"""
    parts = re.split(r"(?<!\\)\|", line.strip())
    if parts and parts[0] == "":
        parts = parts[1:]
    if parts and parts[-1] == "":
        parts = parts[:-1]
    return [part.replace("\\|", "|").strip() for part in parts]


def is_separator(cells: list[str]) -> bool:
    return all(re.fullmatch(r":?-{2,}:?", cell) for cell in cells if cell)


def add_table(doc: Document, rows: list[list[str]]):
    cols = max(len(row) for row in rows)
    table = doc.add_table(rows=len(rows), cols=cols)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True
    for r_index, row in enumerate(rows):
        for c_index in range(cols):
            text = row[c_index] if c_index < len(row) else ""
            cell = table.cell(r_index, c_index)
            paragraph = cell.paragraphs[0]
            paragraph.paragraph_format.space_after = Pt(1)
            add_inline(paragraph, text, base_size=9,
                       color=INK if r_index else ACCENT)
            if r_index == 0:
                for run in paragraph.runs:
                    run.bold = True
                shading = OxmlElement("w:shd")
                shading.set(qn("w:fill"), "E6F4F1")
                cell._tc.get_or_add_tcPr().append(shading)
    doc.add_paragraph()


def add_code_block(doc: Document, lines: list[str]):
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.left_indent = Cm(0.4)
    paragraph.paragraph_format.space_before = Pt(3)
    paragraph.paragraph_format.space_after = Pt(6)
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), "F1F5F9")
    paragraph._p.get_or_add_pPr().append(shading)
    run = paragraph.add_run("\n".join(lines))
    set_run_font(run, name=MONO_FONT, size=9, color=INK)


def add_image(doc: Document, src: str, alt: str, base_dir: str):
    path = src if os.path.isabs(src) else os.path.join(base_dir, src)
    if not os.path.exists(path):
        paragraph = doc.add_paragraph()
        add_inline(paragraph, f"[缺失图片：{src}]", color=RGBColor(0xB4, 0x53, 0x09))
        return
    doc.add_picture(path, width=Inches(6.2))
    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption = doc.add_paragraph()
    caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_inline(caption, alt, base_size=9, color=GRAY)


def convert(md_path: str, docx_path: str, landscape: bool = False):
    with open(md_path, "r", encoding="utf-8") as f:
        lines = f.read().splitlines()

    doc = setup_document(landscape)
    base_dir = os.path.dirname(md_path)

    index = 0
    in_code = False
    code_lines: list[str] = []
    table_lines: list[str] = []

    def flush_table():
        nonlocal table_lines
        if not table_lines:
            return
        rows = [split_row(line) for line in table_lines]
        rows = [row for row in rows if not is_separator(row)]
        if rows:
            add_table(doc, rows)
        table_lines = []

    while index < len(lines):
        line = lines[index]
        stripped = line.strip()

        if stripped.startswith("```"):
            if in_code:
                add_code_block(doc, code_lines)
                code_lines = []
                in_code = False
            else:
                flush_table()
                in_code = True
            index += 1
            continue

        if in_code:
            code_lines.append(line)
            index += 1
            continue

        if stripped.startswith("|"):
            table_lines.append(stripped)
            index += 1
            continue
        flush_table()

        if not stripped:
            index += 1
            continue

        image_match = re.match(r"!\[(.*?)\]\((.*?)\)", stripped)
        if image_match:
            add_image(doc, image_match.group(2), image_match.group(1), base_dir)
            index += 1
            continue

        if re.fullmatch(r"-{3,}", stripped):
            index += 1
            continue

        heading_match = re.match(r"(#{1,4})\s+(.*)", stripped)
        if heading_match:
            level = len(heading_match.group(1))
            paragraph = doc.add_heading(level=level)
            add_inline(paragraph, heading_match.group(2), base_size=18 - level * 1.8)
            index += 1
            continue

        if stripped.startswith(">"):
            paragraph = doc.add_paragraph()
            paragraph.paragraph_format.left_indent = Cm(0.6)
            add_inline(paragraph, stripped.lstrip("> ").strip(), color=GRAY)
            index += 1
            continue

        list_match = re.match(r"^(\s*)([-*]|\d+\.)\s+(.*)", line)
        if list_match:
            depth = len(list_match.group(1)) // 2
            marker = list_match.group(2)
            text = list_match.group(3)
            style = "List Number" if marker[0].isdigit() else "List Bullet"
            paragraph = doc.add_paragraph(style=style)
            paragraph.paragraph_format.left_indent = Cm(0.75 + depth * 0.6)
            add_inline(paragraph, text)
            index += 1
            continue

        paragraph = doc.add_paragraph()
        add_inline(paragraph, stripped)
        index += 1

    flush_table()
    if in_code and code_lines:
        add_code_block(doc, code_lines)

    doc.core_properties.title = os.path.splitext(os.path.basename(docx_path))[0]
    doc.core_properties.author = "______"
    doc.save(docx_path)
    print(f"已导出：{docx_path}")


def main() -> int:
    for md_name, docx_name, landscape in JOBS:
        md_path = os.path.join(DOCS, md_name)
        if not os.path.exists(md_path):
            print(f"[跳过] 源稿不存在：{md_path}")
            continue
        convert(md_path, os.path.join(DOCS, docx_name), landscape)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
